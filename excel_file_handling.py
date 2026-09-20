from openpyxl import Workbook, load_workbook

workbook = Workbook()
sheet = workbook.active

sheet["A1"] = "ID"
sheet["B1"] = "Name"

sheet.append([101, "Ruchitha"])
sheet.append([102, "Anusha"])

workbook.save("students.xlsx")

workbook = load_workbook("students.xlsx")
sheet = workbook.active

for row in sheet.iter_rows(min_row=2, values_only=True):
    print(row[0], row[1])
